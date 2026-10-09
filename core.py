"""Pure logic (no Streamlit): TradingView signal splitting, AlgoTest import, trade analytics.

Formats were derived from the example files, not assumed:
- TradingView export: 2 rows per trade (Exit row first, then Entry row), `Type` is
  "Entry long|Exit long|Entry short|Exit short", dates "YYYY-MM-DD HH:MM" (the manual
  splits show "DD-MM-YYYY HH:MM", every other value identical).
- AlgoTest export: one *parent* row per trade (Type blank; holds Vix and the trade P/L)
  followed by its *leg* rows (Type/Strike/B-S/Qty/prices; P/L repeated). Direction is not
  a column - it is the file the trade came from (Long file = Sell PE, Short file = Sell CE).
"""
import io
import re
from dataclasses import dataclass, field
from datetime import datetime

import openpyxl
import pandas as pd

SIG_REQUIRED = ["Trade number", "Type", "Date and time", "Price"]
ALGO_REQUIRED = ["Index", "Entry Date", "Entry Time", "Exit Date", "Exit Time", "P/L"]
LEG_COLS = ["Type", "Strike", "B/S", "Qty", "Entry Price", "Exit Price", "Expiry"]
DATE_FMTS = {"DD-MM-YYYY HH:MM": "%d-%m-%Y %H:%M",
             "YYYY-MM-DD HH:MM": "%Y-%m-%d %H:%M"}


# ---------- reading / parsing ----------
def read_table(file, name):
    """Read an uploaded .csv/.xlsx into a DataFrame with clear errors."""
    ext = name.lower().rsplit(".", 1)[-1] if "." in name else ""
    if ext not in ("csv", "xlsx"):
        raise ValueError(f"'{name}': unsupported file type '.{ext}'. Upload .csv or .xlsx.")
    try:
        df = pd.read_csv(file, encoding="utf-8-sig") if ext == "csv" else pd.read_excel(file)
    except Exception as e:  # corrupt / unreadable
        raise ValueError(f"'{name}' could not be read ({type(e).__name__}: {e}).") from e
    df.columns = [str(c).strip() for c in df.columns]
    if df.empty:
        raise ValueError(f"'{name}' contains no data rows.")
    return df


def parse_dt(s):
    """Parse ISO or DD-MM-YYYY datetimes (strings or Excel datetimes); unparseable -> NaT."""
    def one(v):
        if isinstance(v, (pd.Timestamp, datetime)):
            return pd.Timestamp(v)
        v = str(v).strip()
        for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%d-%m-%Y %H:%M:%S", "%d-%m-%Y %H:%M"):
            try:
                return pd.Timestamp(datetime.strptime(v, fmt))
            except ValueError:
                pass
        return pd.NaT
    return pd.to_datetime(s.map(one))


# ---------- Stage 1: TradingView signal splitter ----------
@dataclass
class SplitResult:
    long: pd.DataFrame
    short: pd.DataFrame
    unclassified: pd.DataFrame
    total: int
    warnings: list = field(default_factory=list)

    @property
    def reconciled(self):
        return len(self.long) + len(self.short) + len(self.unclassified) == self.total


def split_signals(df, out_fmt="%d-%m-%Y %H:%M"):
    """Split on the long/short suffix of `Type` (so exits stay with their own direction).

    Source row order and all values are kept; only `Date and time` is re-formatted when
    `out_fmt` is given (None = leave untouched)."""
    missing = [c for c in SIG_REQUIRED if c not in df.columns]
    if missing:
        raise ValueError(f"Missing required column(s): {', '.join(missing)}. "
                         f"Found: {', '.join(df.columns)}")
    df = df.reset_index(drop=True)
    side = (df["Type"].astype(str).str.strip()
            .str.extract(r"^(?:entry|exit)\s+(long|short)$", flags=re.I)[0].str.lower())
    out = df.copy()
    warnings = []
    dt = parse_dt(df["Date and time"])
    if out_fmt:
        out["Date and time"] = dt.dt.strftime(out_fmt).where(dt.notna(), df["Date and time"])
    if dt.isna().any():
        warnings.append(f"{int(dt.isna().sum())} row(s) have an unreadable 'Date and time'; "
                        "their original text was kept unchanged.")
    # each trade number should be one entry + one exit of a single direction
    ok = side.notna()
    g = df[ok].assign(_side=side[ok], _entry=df.loc[ok, "Type"].str.lower().str.startswith("entry"))
    bad = g.groupby("Trade number").agg(n=("_side", "size"), sides=("_side", "nunique"),
                                        e=("_entry", "sum"))
    bad = bad[(bad.n != 2) | (bad.sides != 1) | (bad.e != 1)]
    if len(bad):
        warnings.append(f"{len(bad)} trade number(s) are not exactly one Entry + one Exit of a "
                        f"single direction (e.g. {', '.join(map(str, bad.index[:5]))}).")
    unclassified = df[~ok].copy()
    unclassified["Reason"] = "Type '" + unclassified["Type"].astype(str) + "' is not Entry/Exit long/short"
    return SplitResult(out[side == "long"], out[side == "short"], unclassified, len(df), warnings)


def to_xlsx(df):
    buf = io.BytesIO()
    df.to_excel(buf, index=False)
    return buf.getvalue()


def to_csv(df):
    return df.to_csv(index=False).encode("utf-8-sig")


# ---------- Stage 2: AlgoTest import ----------
def _fmt(v):
    return str(int(v)) if isinstance(v, float) and v.is_integer() else str(v)


def _times(s):
    t = s.astype(str).str.extract(r"(\d{1,2}:\d{2}(?::\d{2})?)")[0]
    return pd.to_timedelta(t.where(t.str.count(":") != 1, t + ":00"), errors="coerce")


def _dates(s):
    return pd.to_datetime(s.astype(str).str.strip().str[:10], format="%Y-%m-%d", errors="coerce")


_ALIASES = {"index": "Index", "trade #": "Index", "trade no": "Index", "trade number": "Index",
            "trade no.": "Index", "sr no": "Index", "#": "Index", "p/l": "P/L", "pnl": "P/L", "p&l": "P/L",
            "net p/l": "P/L", "profit/loss": "P/L", "b/s": "B/S", "vix": "Vix", "qty": "Qty",
            # AlgoTest's newer, more detailed export uses different (often hyphenated) headers for the
            # same fields - alias them onto the same canonical names so load_algotest() needs no changes.
            "entry-date": "Entry Date", "entry-time": "Entry Time", "exitdate": "Exit Date",
            "exittime": "Exit Time", "instrument-kind": "Type", "strikeprice": "Strike",
            "position": "B/S", "quantity": "Qty", "expirydate": "Expiry",
            "entry-price": "Entry Price", "exitprice": "Exit Price", "exit-price": "Exit Price",
            **{c.lower(): c for c in ALGO_REQUIRED + LEG_COLS}}


def _canon_columns(df):
    """Match headers ignoring case/extra spaces and common aliases (e.g. 'Trade #' -> 'Index')."""
    ren, used = {}, set()
    for c in df.columns:
        k = _ALIASES.get(re.sub(r"\s+", " ", str(c).strip().lower()))
        if k and k not in used and k not in df.columns:
            ren[c] = k
            used.add(k)
    return df.rename(columns=ren)


def load_algotest(df, direction):
    """Return (trades, rejected, warnings). One row per trade (the AlgoTest parent row).
    Files without Type/leg rows (e.g. a merged trade report) are read one row = one trade."""
    df = _canon_columns(df)
    missing = [c for c in ALGO_REQUIRED if c not in df.columns]
    if missing:
        raise ValueError(f"Not an AlgoTest backtest export - missing column(s): {', '.join(missing)}. "
                         f"Found: {', '.join(df.columns)}")
    df = df.reset_index(drop=True)
    warnings = []
    is_parent = df["Type"].isna() | (df["Type"].astype(str).str.strip() == "") if "Type" in df else \
        pd.Series(True, index=df.index)
    pid = is_parent.cumsum()
    legs = df[~is_parent & (pid > 0)]
    parents = df[is_parent].copy()
    keep_cols = [c for c in parents.columns if c not in LEG_COLS]
    t = parents[keep_cols].copy()
    t["Direction"] = direction
    t["_pid"] = pid[is_parent]
    leg_cols = [c for c in LEG_COLS if c in df.columns]
    if len(legs) and leg_cols:
        agg = legs[leg_cols].apply(lambda c: c.map(_fmt)).groupby(pid[legs.index]).agg(" | ".join)
        t = t.join(agg, on="_pid")
        t["Legs"] = t["_pid"].map(legs.groupby(pid[legs.index]).size()).fillna(0).astype(int)
        # parent P/L is authoritative; cross-check against the sum of its legs
        lp = pd.to_numeric(legs["P/L"], errors="coerce").groupby(pid[legs.index]).sum()
        diff = (pd.to_numeric(t["P/L"], errors="coerce") - t["_pid"].map(lp)).abs()
        if (diff > 1.0).any():
            warnings.append(f"{int((diff > 1.0).sum())} trade(s) where parent P/L differs from the "
                            "sum of its legs by more than 1.0 (parent P/L is used).")
    t = t.drop(columns="_pid")
    t["Entry DateTime"] = _dates(t["Entry Date"]) + _times(t["Entry Time"])
    t["Exit DateTime"] = _dates(t["Exit Date"]) + _times(t["Exit Time"])
    t["P/L"] = pd.to_numeric(t["P/L"], errors="coerce")
    if "Vix" in t:
        t["Vix"] = pd.to_numeric(t["Vix"], errors="coerce")
    t["Duration (min)"] = (t["Exit DateTime"] - t["Entry DateTime"]).dt.total_seconds() / 60

    reason = pd.Series("", index=t.index)
    for cond, msg in ((t["Entry DateTime"].isna(), "invalid entry date/time; "),
                      (t["Exit DateTime"].isna(), "invalid exit date/time; "),
                      (t["P/L"].isna(), "invalid P/L; "),
                      (t["Duration (min)"] < 0, "exit before entry; ")):
        reason = reason.where(~cond, reason + msg)
    dup = t.duplicated(["Entry DateTime", "Exit DateTime", "P/L"], keep="first") & (reason == "")
    reason = reason.where(~dup, "duplicate trade; ")
    rejected = t[reason != ""].assign(Reason=reason[reason != ""].str.rstrip("; "))
    if len(legs) < len(df) - len(parents):
        warnings.append("Some leg rows appear before any trade header row and were ignored.")
    if len(rejected):
        warnings.append(f"{len(rejected)} trade row(s) rejected (see rejected-rows table).")
    return t[reason == ""].reset_index(drop=True), rejected.reset_index(drop=True), warnings


def combine(long_t, short_t):
    """Merge, chronological by entry (then exit). Never merges away rows."""
    parts = [x for x in (long_t, short_t) if x is not None and len(x)]
    if not parts:
        return pd.DataFrame()
    c = pd.concat(parts, ignore_index=True, sort=False)
    c = c.sort_values(["Entry DateTime", "Exit DateTime"], kind="stable").reset_index(drop=True)
    c.insert(0, "Trade #", range(1, len(c) + 1))
    return c


# ---------- Stage 3: analytics ----------
def equity(t):
    """Realised equity curve: trades ordered by exit time (ties: entry time, then file order).
    Cum P&L starts at 0; drawdown = running peak (floored at the 0 baseline) - Cum P&L."""
    e = t.sort_values(["Exit DateTime", "Entry DateTime"], kind="stable").copy()
    e["Cum P&L"] = e["P/L"].cumsum()
    e["Drawdown"] = e["Cum P&L"].cummax().clip(lower=0) - e["Cum P&L"]
    return e


def _streaks(pl):
    w = l = mw = ml = 0
    for x in pl:
        w, l = (w + 1, 0) if x > 0 else (0, l + 1) if x < 0 else (0, 0)
        mw, ml = max(mw, w), max(ml, l)
    return mw, ml


def metrics(t):
    """All KPIs from trade records. None = not computable from the data (never guessed)."""
    n = len(t)
    pl = t["P/L"]
    w, l = pl[pl > 0], pl[pl < 0]
    gp, gl = w.sum(), l.sum()
    eq = equity(t) if n else None
    years = max(0.25, (t["Exit DateTime"].max() - t["Entry DateTime"].min()).days / 365.25) if n else None
    mdd = eq["Drawdown"].max() if n else None
    dd_from = dd_to = None
    if n and mdd > 0:
        i = eq["Drawdown"].values.argmax()
        run = eq["Cum P&L"].cummax().values
        j = max((k for k in range(i + 1) if eq["Cum P&L"].values[k] == run[k] and run[k] >= 0), default=None)
        dd_from = eq["Exit DateTime"].iloc[j] if j is not None else t["Entry DateTime"].min()
        dd_to = eq["Exit DateTime"].iloc[i]
    streaks = _streaks(eq["P/L"]) if n else (0, 0)
    return {
        "Total Trades": n,
        "Winning Trades": len(w),
        "Losing Trades": len(l),
        "Breakeven Trades": n - len(w) - len(l),
        "Win Rate %": len(w) / n * 100 if n else None,        # wins / all trades
        "Gross Profit": gp if n else None,
        "Gross Loss": gl if n else None,
        "Net P&L": pl.sum() if n else None,
        "Avg P&L / Trade": pl.sum() / n if n else None,
        "Avg Win": w.mean() if len(w) else None,
        "Avg Loss": l.mean() if len(l) else None,
        "Profit Factor": gp / -gl if gl < 0 else None,       # undefined with no losing trades
        "Max Drawdown": mdd,
        "Annualised P&L": pl.sum() / years if n else None,      # net / elapsed years (min 0.25y)
        "Avg Monthly P&L": pl.sum() / (years * 12) if n else None,
        "Return / MDD": pl.sum() / years / mdd if n and mdd else None,
        "Max Win Streak": streaks[0], "Max Loss Streak": streaks[1],
        "DD From": dd_from, "DD To": dd_to,
        "Avg Duration (min)": t["Duration (min)"].mean() if n else None,
        "Largest Win": w.max() if len(w) else None,
        "Largest Loss": l.min() if len(l) else None,
    }


def period_pnl(t, freq):
    """Net P&L and trade count per exit-date period ('M' or 'Y')."""
    if t.empty:
        return pd.DataFrame(columns=["Period", "Net P&L", "Trades"])
    g = t.groupby(t["Exit DateTime"].dt.to_period(freq))["P/L"].agg(["sum", "size"])
    return g.rename(columns={"sum": "Net P&L", "size": "Trades"}).rename_axis("Period").reset_index() \
            .assign(Period=lambda d: d["Period"].astype(str))


def monthly_matrix(t):
    """Year x month net P&L by exit date. Returns (matrix with NaN = no trades, yearly metrics frame)."""
    if t.empty:
        return pd.DataFrame(), pd.DataFrame()
    x = t.assign(Y=t["Exit DateTime"].dt.year, M=t["Exit DateTime"].dt.month)
    mat = x.pivot_table(index="Y", columns="M", values="P/L", aggfunc="sum").reindex(columns=range(1, 13))
    rows = []
    for y, g in x.groupby("Y"):
        m = metrics(g)
        rows.append({"Year": y, "Net P&L": m["Net P&L"], "Trades": m["Total Trades"], "Win Rate %": m["Win Rate %"],
                     "Profit Factor": m["Profit Factor"], "Avg P&L / Trade": m["Avg P&L / Trade"],
                     "Max Drawdown": m["Max Drawdown"]})
    return mat, pd.DataFrame(rows)


def rolling_returns(t, months_list=(1, 3, 6, 12)):
    """Overlapping calendar-month windows anchored on every trade; only windows that fit
    entirely inside the sample count. Window P&L = sum of P/L of trades exited in [start, start+N months)."""
    import numpy as np
    out = []
    if t.empty:
        return pd.DataFrame(out)
    e = equity(t)
    d = e["Exit DateTime"].dt.normalize()
    p = np.r_[0.0, e["P/L"].cumsum().values]
    for m in months_list:
        stop = d + pd.DateOffset(months=m)
        ok = (stop <= d.iloc[-1]).values
        if not ok.any():
            continue
        k = d.values.searchsorted(stop[ok].values, side="left")
        w = p[k] - p[np.flatnonzero(ok)]
        out.append({"Months": m, "Windows": len(w), "Worst": w.min(), "Median": float(np.median(w)),
                    "Average": w.mean(), "Best": w.max(), "Positive %": (w > 0).mean() * 100})
    return pd.DataFrame(out)


def to_xlsx_sheets(sheets):
    buf = io.BytesIO()
    with pd.ExcelWriter(buf) as xw:
        for name, df in sheets.items():
            df.to_excel(xw, sheet_name=name[:31], index=False)
    return buf.getvalue()


# ---------- capital (single formula used by the screen AND the Excel export) ----------
def capital_metrics(m, cap):
    """ROI figures from a user-supplied capital/margin. All None when no capital is given.
    Simple (non-compounded) returns: P&L figure / capital x 100. P&L itself never depends on capital."""
    ok = bool(cap) and cap > 0 and m["Total Trades"]
    f = lambda k: (m[k] / cap * 100) if ok and m[k] is not None else None
    return {"ROI %": f("Net P&L"), "Annualised ROI %": f("Annualised P&L"),
            "Avg Monthly ROI %": f("Avg Monthly P&L"), "Max DD % of Capital": f("Max Drawdown")}


# ---------- formatted Excel report ----------
MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
_SUMMARY = [  # (label, metrics key, format)
    ("Net P&L", "Net P&L", "money"), ("Total trades", "Total Trades", "int"),
    ("Winning trades", "Winning Trades", "int"), ("Losing trades", "Losing Trades", "int"),
    ("Breakeven trades", "Breakeven Trades", "int"), ("Win rate %", "Win Rate %", "pct"),
    ("Profit factor", "Profit Factor", "num"), ("Gross profit", "Gross Profit", "money"),
    ("Gross loss", "Gross Loss", "money"), ("Avg P&L per trade", "Avg P&L / Trade", "money"),
    ("Avg winning trade", "Avg Win", "money"), ("Avg losing trade", "Avg Loss", "money"),
    ("Largest win", "Largest Win", "money"), ("Largest loss", "Largest Loss", "money"),
    ("Maximum drawdown", "Max Drawdown", "money"), ("Drawdown from", "DD From", "date"),
    ("Drawdown to", "DD To", "date"), ("Annualised P&L", "Annualised P&L", "money"),
    ("Avg monthly P&L", "Avg Monthly P&L", "money"), ("Return / MDD (annualised)", "Return / MDD", "num"),
    ("Avg trade duration (min)", "Avg Duration (min)", "int"), ("Longest win streak", "Max Win Streak", "int"),
    ("Longest loss streak", "Max Loss Streak", "int"),
]
_CAP = [("ROI on capital %", "ROI %", "pct"), ("Annualised ROI % (simple)", "Annualised ROI %", "pct"),
        ("Avg monthly ROI %", "Avg Monthly ROI %", "pct"), ("Max drawdown % of capital", "Max DD % of Capital", "pct")]
_FMT = {
    **{c: "money" for c in ("Net P&L", "Avg P&L / Trade", "Avg P&L per trade", "Max Drawdown", "Max drawdown", "Worst",
                            "Median", "Average", "Best", "Worst trade", "Annualised P&L", "Total", "Max DD", "P/L")},
    **{c: "pct" for c in ("Win Rate %", "Win rate %", "ROI %", "Positive windows %", "Max DD % of capital",
                          "Annualised ROI %", "Worst % of capital", "Average % of capital", "Best % of capital")},
    **{c: "int" for c in ("Trades", "Windows", "Holding period (months)", "Longest losing run", "Trade #",
                          "Duration (min)", "Legs")},
    **{c: "num" for c in ("Profit Factor", "Profit factor", "Return / MDD", "Vix")},
    "Entry DateTime": "dt", "Exit DateTime": "dt", "Year": "year",
}


def report_tables(series, cap=None, split=None, roll_months=(1, 3, 6, 12), trades=None, rejected=None):
    """Every report table as a DataFrame: {sheet name: (df, {column: format}, extras)}.
    `series` = {"Combined": df, "Long": df, "Short": df}. Formats: money, pct, int, num, date, dt, text."""
    names = list(series)
    M = {n: metrics(s) for n, s in series.items()}
    C = {n: capital_metrics(M[n], cap) for n in names}
    out = {}

    spec = _SUMMARY[:1] + (_CAP if cap else []) + _SUMMARY[1:]   # all capital rows sit right under Net P&L
    rows = [[label] + [(C if k in C[n] else M)[n][k] for n in names] for label, k, _ in spec]
    out["Summary"] = (pd.DataFrame(rows, columns=["Metric"] + names), {},
                      {"row_formats": [f for _, _, f in spec], "note": "One column per series; each computed from its own trades."})

    yr, roll, risk, reg = [], [], [], []
    for n in names:
        s, m = series[n], M[n]
        mat, y = monthly_matrix(s)
        if len(y):
            y = y.copy()
            if cap:
                y["ROI %"] = y["Net P&L"] / cap * 100
            yr.append(y.assign(Series=n))
            mt = mat.copy()
            mt.columns = MONTHS
            yi = y.set_index("Year")
            mt["Total"], mt["Max DD"] = yi["Net P&L"], yi["Max Drawdown"]
            if cap:
                mt["ROI %"] = mt["Total"] / cap * 100
            mt.index = mt.index.astype(str)
            allrow = {**mat.sum(min_count=1).set_axis(MONTHS).to_dict(), "Total": m["Net P&L"], "Max DD": m["Max Drawdown"]}
            if cap:
                allrow["ROI %"] = C[n]["ROI %"]
            mt.loc["All years"] = pd.Series(allrow)
            out[f"Monthly {n}"] = (mt.rename_axis("Year").reset_index(), {**{k: "money" for k in MONTHS}, **_FMT, "Year": "text"},
                                   {"heat": MONTHS, "note": f"{n}: net P&L by exit month (blank = no trades). "
                                                            "Max DD is within each year; 'All years' row uses the full period."})
        r = rolling_returns(s, tuple(roll_months))
        if len(r):
            r = r.rename(columns={"Months": "Holding period (months)", "Positive %": "Positive windows %"})
            if cap:
                for col in ("Worst", "Average", "Best"):
                    r[f"{col} % of capital"] = r[col] / cap * 100
            roll.append(r.assign(Series=n))
        risk.append({"Series": n, "Max drawdown": m["Max Drawdown"], "Worst trade": m["Largest Loss"],
                     "Longest losing run": m["Max Loss Streak"], "Annualised P&L": m["Annualised P&L"],
                     "Return / MDD": m["Return / MDD"],
                     **({"Max DD % of capital": C[n]["Max DD % of Capital"], "Annualised ROI %": C[n]["Annualised ROI %"]}
                        if cap else {})})
        if split is not None:
            for tag, sub in (("Full period", s), (f"From {pd.Timestamp(split):%d %b %Y}", s[s["Exit DateTime"] >= split])):
                if sub.empty:
                    continue
                mm = metrics(sub)
                reg.append({"Series": n, "Period": tag, "Net P&L": mm["Net P&L"], "Trades": mm["Total Trades"],
                            "Win rate %": mm["Win Rate %"], "Profit factor": mm["Profit Factor"],
                            "Avg P&L per trade": mm["Avg P&L / Trade"], "Max drawdown": mm["Max Drawdown"],
                            "Return / MDD": mm["Return / MDD"], **({"ROI %": mm["Net P&L"] / cap * 100} if cap else {})})

    lead = lambda df: df[["Series"] + [c for c in df.columns if c != "Series"]]
    if yr:
        out["Yearly"] = (lead(pd.concat(yr, ignore_index=True)), _FMT, {"pnl": ["Net P&L"], "note": "Calendar years by exit date."})
    if reg:
        out["Regime"] = (pd.DataFrame(reg), _FMT, {"pnl": ["Net P&L"], "note": "Full period vs. the period from the split date."})
    if roll:
        out["Rolling"] = (lead(pd.concat(roll, ignore_index=True)), _FMT,
                          {"pnl": ["Worst", "Median", "Average", "Best"],
                           "note": "Overlapping calendar windows anchored on every trade; only windows that fit in the data."})
    out["Risk"] = (pd.DataFrame(risk), _FMT, {"note": "Closed-trade drawdown; Return / MDD = annualised P&L / max drawdown."})
    if trades is not None and len(trades):
        out["Trades"] = (trades, {**_FMT, "Index": "text"}, {"pnl": ["P/L"], "note": "Every trade in the report (after filters)."})
    if rejected is not None and len(rejected):
        out["Rejected rows"] = (rejected, _FMT, {"note": "Rows excluded from every figure, with the reason."})
    order = ["Summary", "Yearly"] + [k for k in out if k.startswith("Monthly")] + ["Regime", "Rolling", "Risk", "Trades", "Rejected rows"]
    return {k: out[k] for k in order if k in out}


_XL = {"money": '"₹"#,##0.00;[Red]-"₹"#,##0.00', "pct": '0.00"%";[Red]-0.00"%"', "int": "#,##0", "year": "0",
       "num": "0.00", "date": "dd-mmm-yyyy", "dt": "dd-mmm-yyyy hh:mm", "text": "@"}


def report_workbook(info, tables, title="Long + Short Backtest Report"):
    """Formatted .xlsx: a 'Report Info' sheet, then one styled Excel Table per report table
    (navy header, banded rows, sort/filter, number formats, green/red P&L, heatmap on monthly sheets)."""
    from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo

    line = Side(style="thin", color="D5DAE0")
    box = Border(top=line, bottom=line, left=line, right=line)
    head_font, head_fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F4E79")
    green, red = Font(color="1E7B45", bold=True), Font(color="B3261E", bold=True)

    def sign_rules(ws, rng):
        ws.conditional_formatting.add(rng, CellIsRule(operator="greaterThan", formula=["0"], font=green))
        ws.conditional_formatting.add(rng, CellIsRule(operator="lessThan", formula=["0"], font=red))

    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        pd.DataFrame(info, columns=["Item", "Value"]).to_excel(xw, sheet_name="Report Info", index=False, startrow=3)
        for name, (df, _, _) in tables.items():
            df.to_excel(xw, sheet_name=name[:31], index=False, startrow=3)
        wb = xw.book

        ws = wb["Report Info"]
        ws["A1"], ws["A2"] = title, "Built from the AlgoTest trade files. ₹ values are as exported; % figures need a capital."
        ws["A1"].font, ws["A2"].font = Font(bold=True, size=16, color="16324F"), Font(italic=True, color="5A6675")
        ws.sheet_view.showGridLines = False
        for c in ws[4]:
            c.font, c.fill, c.border = head_font, head_fill, box
        for r in ws.iter_rows(min_row=5):
            r[0].font, r[0].border, r[1].border = Font(bold=True, color="25313F"), box, box
            r[1].alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 30, 110

        for i, (name, (df, fmts, extra)) in enumerate(tables.items()):
            ws = wb[name[:31]]
            ws["A1"], ws["A2"] = name, extra.get("note", "")
            ws["A1"].font, ws["A2"].font = Font(bold=True, size=14, color="16324F"), Font(italic=True, color="5A6675")
            ws.freeze_panes = "B5" if name.startswith(("Summary", "Monthly")) else "A5"
            ws.sheet_view.showGridLines = False
            ws.page_setup.orientation, ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = "landscape", 1, 0
            ws.sheet_properties.pageSetUpPr.fitToPage = True
            ws.print_title_rows = "4:4"
            if df.empty:
                ws["A4"] = "No data for the current filters."
                continue
            nr, nc = len(df) + 4, len(df.columns)
            t = Table(displayName=f"T{i}_" + re.sub(r"\W", "_", name), ref=f"A4:{get_column_letter(nc)}{nr}")
            t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
            ws.add_table(t)
            for c in ws[4]:
                c.font, c.fill, c.border = head_font, head_fill, box
                c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            ws.row_dimensions[4].height = 32
            rowf = extra.get("row_formats")
            for j, col in enumerate(df.columns, 1):
                L = get_column_letter(j)
                for r in range(5, nr + 1):
                    cell = ws[f"{L}{r}"]
                    cell.border = box
                    fmt = _XL.get(rowf[r - 5] if rowf and j > 1 else fmts.get(col, ""))
                    if fmt:
                        cell.number_format = fmt
                    if rowf and j > 1 and cell.value is None:
                        cell.value = "n/a"
                        cell.alignment = Alignment(horizontal="right")
                    if j == 1:
                        cell.font = Font(bold=True, color="25313F")
                vals = df[col].head(500)
                w = max([len(str(col)) + 2] + [len(f"{x:,.2f}") + 3 if isinstance(x, float) else len(str(x)) for x in vals])
                ws.column_dimensions[L].width = min(max(w, 10), 48)
            for col in extra.get("pnl", []):
                if col in df.columns:
                    L = get_column_letter(list(df.columns).index(col) + 1)
                    sign_rules(ws, f"{L}5:{L}{nr}")
            heat = [c for c in extra.get("heat", []) if c in df.columns]
            if heat:
                a = get_column_letter(list(df.columns).index(heat[0]) + 1)
                b = get_column_letter(list(df.columns).index(heat[-1]) + 1)
                ws.conditional_formatting.add(f"{a}5:{b}{nr - 1}", ColorScaleRule(
                    start_type="min", start_color="F4A6A1", mid_type="num", mid_value=0, mid_color="FFFFFF",
                    end_type="max", end_color="9FD5B5"))
                for col in ("Total", "ROI %"):
                    if col in df.columns:
                        L = get_column_letter(list(df.columns).index(col) + 1)
                        sign_rules(ws, f"{L}5:{L}{nr}")
                for c in ws[nr]:  # 'All years' total row
                    c.font = Font(bold=True, color="25313F")
                    c.border = Border(top=Side(style="medium", color="16324F"), bottom=line, left=line, right=line)
            if rowf:  # Summary: colour P&L-type rows by sign
                for r, f in enumerate(rowf, 5):
                    if f in ("money", "pct") and not str(ws[f"A{r}"].value).lower().startswith(("gross loss", "max", "largest loss", "avg losing", "win rate")):
                        sign_rules(ws, f"B{r}:{get_column_letter(nc)}{r}")
    return buf.getvalue()


# ---------- Strategy Hub: merge many AlgoTest / StockMock files into one trade log ----------
def is_stockmock_workbook(file_bytes):
    """Quick format sniff on an .xlsx's sheet names, no full parse."""
    try:
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True)
        sheets = wb.sheetnames
        wb.close()
    except Exception:
        return False
    return "Basket Strategies" in sheets


def load_stockmock(file_bytes, name):
    """Return (trades, rejected, warnings). One row per expiry/trade-cycle, read from each
    enabled (Run=True) strategy's '# S-n - Result' sheet. StockMock gives only the net P/L per
    cycle (no per-leg strike/price detail) and only an EXIT time (no entry time) - Duration is
    therefore left blank for these rows rather than guessed."""
    try:
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    except Exception as e:
        raise ValueError(f"'{name}' could not be read as an Excel workbook ({type(e).__name__}: {e}).") from e
    if "Basket Strategies" not in wb.sheetnames:
        raise ValueError(f"'{name}' does not look like a StockMock basket export "
                         "(no 'Basket Strategies' sheet).")

    bs_rows = list(wb["Basket Strategies"].iter_rows(values_only=True))
    hdr = next((i for i, r in enumerate(bs_rows) if r and r[0] == "Run"), None)
    active = {}  # "#S-1" -> strategy name
    if hdr is not None:
        for r in bs_rows[hdr + 1:]:
            if r and r[0] and str(r[0]).strip().lower() == "true" and r[1]:
                active[str(r[1]).strip()] = str(r[2] or r[1]).strip()
    if not active:
        raise ValueError(f"'{name}': no enabled (Run=True) strategy found in 'Basket Strategies'.")

    norm = lambda s: re.sub(r"\s+", "", s).lower()
    rows, warnings = [], []
    for sid, nm in active.items():
        sheet = next((s for s in wb.sheetnames if norm(s) == norm(f"{sid} - Result")), None)
        if not sheet:
            warnings.append(f"'{nm}' ({sid}) is enabled but its '- Result' sheet was not found; skipped.")
            continue
        rws = list(wb[sheet].iter_rows(values_only=True))
        hi = next((i for i, r in enumerate(rws) if r and r[0] == "Include"), None)
        if hi is None:
            warnings.append(f"'{nm}' ({sid}): result sheet layout not recognised; skipped.")
            continue
        ci = {str(c).strip(): j for j, c in enumerate(rws[hi]) if c}
        missing = [k for k in ("Expiry", "Entry Date", "Exit Time", "Profit") if k not in ci]
        if missing:
            warnings.append(f"'{nm}' ({sid}): missing column(s) {missing}; skipped.")
            continue
        for r in rws[hi + 1:]:
            if not r or r[0] is None:
                continue
            rows.append({"Strategy": nm, "Expiry": r[ci["Expiry"]], "_entry": r[ci["Entry Date"]],
                        "_exit": r[ci["Exit Time"]], "P/L": r[ci["Profit"]]})
    if not rows:
        raise ValueError(f"'{name}': none of the enabled strategies had a readable result sheet.")

    t = pd.DataFrame(rows)
    t["Entry DateTime"] = pd.to_datetime(t["_entry"].astype(str).str.extract(r"(\d{4}-\d{2}-\d{2})")[0],
                                         errors="coerce")
    ex = t["_exit"].astype(str)
    t["Exit DateTime"] = (pd.to_datetime(ex.str.extract(r"(\d{4}-\d{2}-\d{2})")[0], errors="coerce")
                          + pd.to_timedelta(ex.str.extract(r"(\d{1,2}:\d{2})\s*$")[0] + ":00", errors="coerce"))
    t["P/L"] = pd.to_numeric(t["P/L"], errors="coerce")
    t["Index"] = [f"SM-{i + 1}" for i in range(len(t))]
    t["Remarks"] = "Strategy: " + t["Strategy"]
    t["Duration (min)"] = pd.NA  # no entry time in the source - never fabricated
    t = t.drop(columns=["_entry", "_exit", "Strategy"])

    reason = pd.Series("", index=t.index)
    for cond, msg in ((t["Entry DateTime"].isna(), "invalid entry date; "),
                      (t["Exit DateTime"].isna(), "invalid exit date/time; "),
                      (t["P/L"].isna(), "invalid P/L; ")):
        reason = reason.where(~cond, reason + msg)
    rejected = t[reason != ""].assign(Reason=reason[reason != ""].str.rstrip("; "))
    t = t[reason == ""].reset_index(drop=True)
    if len(rejected):
        warnings.append(f"{len(rejected)} row(s) rejected (see rejected-rows table).")
    return t, rejected.reset_index(drop=True), warnings


def detect_and_load(file_bytes, name):
    """Auto-detect an uploaded file as a StockMock basket workbook or an AlgoTest export
    (either column layout) and parse it. Returns (trades, rejected, warnings, kind)."""
    ext = name.lower().rsplit(".", 1)[-1] if "." in name else ""
    if ext == "xlsx" and is_stockmock_workbook(file_bytes):
        t, rej, warn = load_stockmock(file_bytes, name)
        return t, rej, warn, "StockMock"
    df = read_table(io.BytesIO(file_bytes), name)
    t, rej, warn = load_algotest(df, name)
    return t, rej, warn, "AlgoTest"


def hub_merge(frames, capital, charges):
    """Merge parsed trade frames into one chronological master log. Total charges are spread
    evenly across every trade (the source files carry no per-trade charge). `P/L` is replaced
    with Net P/L so metrics()/equity()/monthly_matrix() work on it unchanged; Gross P/L and
    Charges stay as their own columns."""
    parts = [f for f in frames if f is not None and len(f)]
    if not parts:
        return pd.DataFrame()
    m = pd.concat(parts, ignore_index=True, sort=False)
    m = m.sort_values(["Exit DateTime", "Entry DateTime"], kind="stable").reset_index(drop=True)
    n = len(m)
    charge = (charges / n) if charges and n else 0.0
    m["Gross P/L"] = m["P/L"]
    m["Charges"] = charge
    m["Net P/L"] = m["Gross P/L"] - m["Charges"]
    m["Return %"] = (m["Net P/L"] / capital * 100) if capital else None
    m["P/L"] = m["Net P/L"]
    m.insert(0, "Trade #", range(1, n + 1))
    return m


def hub_equity(t, capital):
    """Equity curve starting at `capital` (not 0), with the standard peak-relative drawdown %
    (underwater curve) - distinct from Max-Drawdown-as-%-of-capital, which hub_metrics gives."""
    e = equity(t)
    e["Equity"] = capital + e["Cum P&L"]
    peak = e["Equity"].cummax()
    e["Drawdown %"] = ((peak - e["Equity"]) / peak.where(peak != 0)).fillna(0) * 100
    return e


def hub_metrics(t, capital):
    """metrics() plus capital-based figures: ROI % and Max Drawdown as a simple % of capital
    (not peak-relative - that distinction matters, see hub_equity)."""
    m = metrics(t)
    ok = bool(capital) and m["Total Trades"]
    m["ROI %"] = (m["Net P&L"] / capital * 100) if ok else None
    m["Max Drawdown % of Capital"] = (m["Max Drawdown"] / capital * 100) if ok and m["Max Drawdown"] else None
    return m


def hub_monthly_matrix(t, capital):
    """monthly_matrix() with each year's Return % against the fixed capital (not annualised)."""
    mat, yearly = monthly_matrix(t)
    if capital and len(yearly):
        yearly = yearly.copy()
        yearly["Return %"] = yearly["Net P&L"] / capital * 100
    return mat, yearly
