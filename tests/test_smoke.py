"""Self-contained checks (no private data): synthetic TradingView + AlgoTest files run through
the whole pipeline, and every page of the app renders without an exception. Runs on GitHub CI."""
import io
import pathlib
import sys

import openpyxl
import pandas as pd
import pytest
from streamlit.testing.v1 import AppTest

ROOT = pathlib.Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
import core  # noqa: E402


def tradingview(n=40):
    """n trades alternating long/short, exported the TradingView way (exit row, then entry row)."""
    rows, t0 = [], pd.Timestamp("2024-01-01 09:30")
    for i in range(1, n + 1):
        side = "long" if i % 2 else "short"
        e = t0 + pd.Timedelta(days=i)
        pnl = (-1) ** i * 10.5 * i
        for kind, ts in (("Exit", e + pd.Timedelta(hours=2)), ("Entry", e)):
            rows.append({"Trade number": i, "Type": f"{kind} {side}", "Date and time": f"{ts:%Y-%m-%d %H:%M}",
                         "Signal": "entry" if kind == "Entry" else "exit", "Price": 20000 + i, "Size (qty)": 1,
                         "Net PnL NONE": pnl})
    return pd.DataFrame(rows)


def algotest(n, pnl, start="2024-01-01"):
    """AlgoTest layout: a parent row (Type blank, holds P/L) followed by one leg row per trade."""
    rows, d = [], pd.Timestamp(start)
    for i in range(1, n + 1):
        day = d + pd.Timedelta(days=i)
        base = {"Entry Date": f"{day:%Y-%m-%d}", "Entry Time": " 10:00:00", "Exit Date": f"{day:%Y-%m-%d}",
                "Exit Time": " 14:30:00"}
        p = pnl(i)
        rows.append({"Index": str(i), **base, "Type": None, "Strike": None, "B/S": None, "Qty": None,
                     "Entry Price": None, "Exit Price": None, "Vix": 13.5, "P/L": p})
        rows.append({"Index": f"{i}.1", **base, "Type": "PE", "Strike": 22000, "B/S": "Sell", "Qty": 65,
                     "Entry Price": 100.0, "Exit Price": 100 - p / 65, "Vix": None, "P/L": p})
    return pd.DataFrame(rows)


def test_splitter_end_to_end():
    src = tradingview(40)
    r = core.split_signals(src)
    assert r.reconciled and (len(r.long), len(r.short), len(r.unclassified)) == (40, 40, 0) and not r.warnings
    assert set(r.long.Type.str.split().str[1]) == {"long"} and set(r.short.Type.str.split().str[1]) == {"short"}
    assert openpyxl.load_workbook(io.BytesIO(core.to_xlsx(r.long))).active.max_row == 41


def test_report_end_to_end_with_capital():
    L, _, _ = core.load_algotest(algotest(30, lambda i: 100.0 if i % 3 else -150.0), "Long")
    S, _, _ = core.load_algotest(algotest(20, lambda i: -80.0 if i % 4 == 0 else 60.0, "2024-03-01"), "Short")
    C = core.combine(L, S)
    m = core.metrics(C)
    assert m["Total Trades"] == 50 and m["Net P&L"] == pytest.approx(L["P/L"].sum() + S["P/L"].sum())
    assert core.capital_metrics(m, 100000)["ROI %"] == pytest.approx(m["Net P&L"] / 1000)
    tables = core.report_tables({"Combined": C, "Long": L, "Short": S}, 100000, pd.Timestamp("2024-03-15"), (1,), C)
    wb = openpyxl.load_workbook(io.BytesIO(core.report_workbook([("Capital", "100000")], tables)))
    assert {"Report Info", "Summary", "Yearly", "Trades"} <= set(wb.sheetnames)
    assert wb["Trades"].max_row - 4 == 50


@pytest.mark.parametrize("page", ["app.py", "views/0_Instructions.py", "views/1_Signal_Splitter.py",
                                  "views/2_Backtest_Dashboard.py", "views/3_Strategy_Hub.py"])
def test_every_page_renders(page):
    at = AppTest.from_file(str(ROOT / page), default_timeout=60).run()
    assert not at.exception, [e.value for e in at.exception]


# ---------- Strategy Hub: AlgoTest (detailed columns) + StockMock + multi-file merge ----------
def algotest_detailed_csv(n=10, start="2024-10-07"):
    """AlgoTest's newer export: hyphenated/condensed headers, one parent + one leg row per trade."""
    rows, d = [], pd.Timestamp(start)
    for i in range(1, n + 1):
        day = d + pd.Timedelta(days=i)
        pnl = 325.0 if i % 2 else -200.0
        common = {"Entry-Date": f"{day:%Y-%m-%d}", "Entry-Weekday": None, "Entry-Time": "10:40:00",
                 "ExitDate": f"{day:%Y-%m-%d}", "Exit-Weekday": None, "ExitTime": "11:27:00",
                 "P/L": pnl, "P/L-Percentage": None, "Highest MTM(Candle Close)": None,
                 "Lowest MTM(Candle Close)": None, "Remarks": "Stop Loss Hit" if pnl < 0 else "Target Hit"}
        rows.append({"Index": str(i), **common, "Entry-Price": None, "Entry-Delta": None, "Quantity": None,
                    "Instrument-Kind": None, "StrikePrice": None, "Position": None, "ExitPrice": None,
                    "Exit-Delta": None, "ExpiryDate": None})
        rows.append({"Index": f"{i}.1", **common, "Entry-Price": 20.0, "Entry-Delta": 0.1, "Quantity": 65,
                    "Instrument-Kind": "PE", "StrikePrice": 25000.0, "Position": "Sell", "ExitPrice": 15.0,
                    "Exit-Delta": 0.08, "ExpiryDate": f"{(day + pd.Timedelta(days=3)):%Y-%m-%d}"})
    return pd.DataFrame(rows)


def stockmock_workbook(strategies=(("#S-1", "Test_Strategy_A"), ("#S-2", "Test_Strategy_B"))):
    """A minimal in-memory StockMock basket workbook: 'Basket Strategies' + one '# S-n - Result' per strategy."""
    wb = openpyxl.Workbook()
    bs = wb.active
    bs.title = "Basket Strategies"
    bs.append(["Run", "#", "Name"])
    for sid, name in strategies:
        bs.append(["True", sid, name])
    for i, (sid, name) in enumerate(strategies, 1):
        ws = wb.create_sheet(f"{sid} - Result")
        ws.append(["Include", "Expiry", "Date", "Day", "Entry Date", "Exit Time", "Profit"])
        for d in range(1, 4):
            day = pd.Timestamp("2026-01-01") + pd.Timedelta(days=d * 7)
            ws.append(["True", f"{day:%d %b %Y}".upper(), f"{day:%d-%b-%Y}", f"{day:%a}",
                      f"{day:%Y-%m-%d} \n( {day:%a})", f"{day:%Y-%m-%d}({day:%a}) 15:38", 1000 * i + d])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_hub_detailed_algotest_columns_are_recognised():
    """The new hyphenated AlgoTest column layout must parse through the SAME load_algotest()."""
    t, rej, warn, kind = core.detect_and_load(
        algotest_detailed_csv(10).to_csv(index=False).encode("utf-8-sig"), "detailed.csv")
    assert kind == "AlgoTest" and len(t) == 10 and rej.empty
    assert set(t["Type"]) == {"PE"} and set(t["B/S"]) == {"Sell"}
    assert t["Entry Price"].astype(float).eq(20.0).all() and t["Exit Price"].astype(float).eq(15.0).all()
    assert t["Expiry"].notna().all()


def test_hub_stockmock_workbook_parses():
    t, rej, warn, kind = core.detect_and_load(stockmock_workbook(), "basket.xlsx")
    assert kind == "StockMock" and len(t) == 6 and rej.empty  # 2 strategies x 3 cycles
    assert t["Duration (min)"].isna().all()  # no entry time in the source - never fabricated
    assert t["Entry DateTime"].notna().all() and t["Exit DateTime"].notna().all()
    assert core.is_stockmock_workbook(stockmock_workbook())
    not_sm = io.BytesIO()
    openpyxl.Workbook().save(not_sm)  # a plain xlsx with no 'Basket Strategies' sheet
    assert not core.is_stockmock_workbook(not_sm.getvalue())
    with pytest.raises(ValueError, match="Basket Strategies"):
        core.load_stockmock(not_sm.getvalue(), "bad.xlsx")
    with pytest.raises(ValueError, match="could not be read"):
        core.load_stockmock(b"not an xlsx", "corrupt.xlsx")


def test_hub_merge_multi_file_capital_and_charges():
    algo, _, _, _ = core.detect_and_load(algotest_detailed_csv(10).to_csv(index=False).encode("utf-8-sig"), "a.csv")
    sm, _, _, _ = core.detect_and_load(stockmock_workbook(), "b.xlsx")
    m = core.hub_merge([algo, sm], capital=100000, charges=1000)
    assert len(m) == 16
    assert m["Charges"].eq(1000 / 16).all()
    assert (m["Gross P/L"] - m["Charges"] - m["Net P/L"]).abs().max() < 1e-9
    assert m["Return %"].eq(m["Net P/L"] / 100000 * 100).all()
    mm = core.hub_metrics(m, 100000)
    assert mm["ROI %"] == pytest.approx(mm["Net P&L"] / 100000 * 100)
    assert mm["Total Trades"] == 16
    eq = core.hub_equity(m, 100000)
    assert eq["Equity"].iloc[0] == pytest.approx(100000 + m.sort_values("Exit DateTime")["Net P/L"].iloc[0])
    assert (eq["Drawdown %"] >= 0).all()
    mat, yearly = core.hub_monthly_matrix(m, 100000)
    assert "Return %" in yearly.columns


def test_hub_merge_with_no_files_or_zero_charges():
    assert core.hub_merge([], 100000, 0).empty
    algo, _, _, _ = core.detect_and_load(algotest_detailed_csv(4).to_csv(index=False).encode("utf-8-sig"), "a.csv")
    m = core.hub_merge([algo], 100000, 0)
    assert m["Charges"].eq(0).all() and (m["Gross P/L"] == m["Net P/L"]).all()
    m0 = core.hub_metrics(m.iloc[0:0], 100000)
    assert m0["ROI %"] is None  # no trades - never divides by zero into a fake number


def test_hub_rejects_unrecognised_file():
    with pytest.raises(ValueError):
        core.detect_and_load(pd.DataFrame({"A": [1], "B": [2]}).to_csv(index=False).encode(), "junk.csv")


def test_hub_detailed_algotest_accepts_dd_mm_yyyy_dates():
    """Regression: some exports of the detailed AlgoTest format use DD-MM-YYYY (e.g. after a
    round-trip through Excel) instead of AlgoTest's usual YYYY-MM-DD - every row was being
    rejected as 'invalid entry date/time' until _dates() learned the fallback format."""
    df = algotest_detailed_csv(6)
    for c in ("Entry-Date", "ExitDate"):
        df[c] = pd.to_datetime(df[c]).dt.strftime("%d-%m-%Y")
    t, rej, warn, kind = core.detect_and_load(df.to_csv(index=False).encode("utf-8-sig"), "dd-mm-yyyy.csv")
    assert kind == "AlgoTest" and len(t) == 6 and rej.empty
    expected = [f"{(pd.Timestamp('2024-10-07') + pd.Timedelta(days=i)):%Y-%m-%d}" for i in range(1, 7)]
    assert t.sort_values("Entry DateTime")["Entry DateTime"].dt.strftime("%Y-%m-%d").tolist() == expected


def own_report_workbook(n=10, with_source_col=False):
    """A minimal stand-in for this app's own 'Download full report' export: a 'Report Info'
    sheet plus a 'Trades' sheet with the header on row 4 (matches core.report_workbook())."""
    import core as _core
    cols = ["Trade #", "Source" if with_source_col else "Direction", "Entry DateTime", "Exit DateTime",
           "Duration (min)", "P/L"]
    rows = []
    for i in range(1, n + 1):
        e = pd.Timestamp("2025-01-01") + pd.Timedelta(days=i)
        rows.append([i, "prior.csv" if with_source_col else ("Long" if i % 2 else "Short"),
                    e, e + pd.Timedelta(hours=2), 120, 100.0 if i % 2 else -50.0])
    trades = pd.DataFrame(rows, columns=cols)
    info = pd.DataFrame({"Item": ["x"], "Value": ["y"]})
    buf = io.BytesIO()
    with pd.ExcelWriter(buf) as xw:
        info.to_excel(xw, sheet_name="Report Info", index=False, startrow=3)
        trades.to_excel(xw, sheet_name="Trades", index=False, startrow=3)
    return buf.getvalue()


def test_hub_reimports_its_own_previous_report():
    assert core.is_own_report_workbook(own_report_workbook())
    t, rej, warn, kind = core.detect_and_load(own_report_workbook(10), "old.xlsx")
    assert kind == "Previous report" and len(t) == 10 and rej.empty
    assert "Direction" not in t.columns  # folded into Remarks, not left to collide with Source
    assert t["Remarks"].str.contains("Re-imported").all()
    assert pd.api.types.is_datetime64_any_dtype(t["Entry DateTime"])

    # re-importing a PREVIOUS HUB export (which already has its own Source/Gross P/L/Charges/
    # Net P/L columns, not Direction) must not collide with hub_merge()'s own column names
    t2, _, _, _ = core.detect_and_load(own_report_workbook(4, with_source_col=True), "old_hub.xlsx")
    assert "Source" in t2.columns and len(t2) == 4


def test_hub_merge_with_reimported_report_does_not_duplicate_trade_number():
    """Regression: a re-imported report's 'Trades' sheet already has a 'Trade #' column;
    hub_merge() used to crash with 'cannot insert Trade #, already exists'."""
    algo, _, _, _ = core.detect_and_load(algotest_detailed_csv(5).to_csv(index=False).encode("utf-8-sig"), "a.csv")
    old, _, _, _ = core.detect_and_load(own_report_workbook(5), "old.xlsx")
    m = core.hub_merge([algo, old], capital=100000, charges=0)
    assert len(m) == 10
    assert m["Trade #"].tolist() == list(range(1, 11))
    assert not m["Trade #"].duplicated().any()


def test_hub_detect_and_load_never_leaks_a_raw_exception(monkeypatch):
    """Regression: Streamlit Cloud reported an uncaught AttributeError from inside
    detect_and_load(), which crashed the whole page instead of being shown as a per-file
    error. Whatever goes wrong inside parsing, callers must only ever see a ValueError."""
    def boom(*a, **k):
        raise AttributeError("something unexpected")
    monkeypatch.setattr(core, "load_stockmock", boom)
    with pytest.raises(ValueError, match="unexpected AttributeError"):
        core.detect_and_load(stockmock_workbook(), "basket.xlsx")

    monkeypatch.setattr(core, "load_algotest", boom)
    with pytest.raises(ValueError, match="unexpected AttributeError"):
        core.detect_and_load(algotest_detailed_csv(2).to_csv(index=False).encode("utf-8-sig"), "a.csv")


def test_hub_stockmock_only_trade_log_has_no_strike_column():
    """Regression: a StockMock-only upload has no 'Strike'/'Type' column at all (AlgoTest's
    leg-join never runs), and the trade log page crashed with KeyError on disp["Strike"].
    The page now creates these columns defensively before building Strike/Type - replicate
    that same guard here against core's raw output."""
    t, _, _, _ = core.detect_and_load(stockmock_workbook(), "basket.xlsx")
    assert "Strike" not in t.columns and "Type" not in t.columns
    m = core.hub_merge([t], capital=100000, charges=0)
    disp = m.copy()
    for c in ("Strike", "Type", "Entry Price", "Exit Price", "B/S", "Expiry", "Remarks"):
        if c not in disp:
            disp[c] = None
    disp["Strike/Type"] = (disp["Strike"].map(lambda x: "" if pd.isna(x) else str(x)) + " "
                           + disp["Type"].fillna("")).str.strip().replace("", "—")
    assert (disp["Strike/Type"] == "—").all()


def test_stockmock_margin_and_daily_pnl_match_stockmocks_own_figures():
    """A StockMock workbook can print its own basket-level 'Estimated Margin' in the
    'Basket Strategies' sheet - extract it exactly, and verify daily_pnl() groups trades the
    same way StockMock itself does for its day-level Win% figures (used for the day-wise
    breakdown table and to let users cross-check a Hub import against the source file)."""
    wb = stockmock_workbook()
    assert core.extract_stockmock_margin(wb) is None  # the minimal test fixture has no RESULT block

    import openpyxl
    book = openpyxl.load_workbook(io.BytesIO(wb))
    bs = book["Basket Strategies"]
    bs.append(["Estimated Margin (On tue)", None, None, "Rs 5.25L"])
    buf = io.BytesIO()
    book.save(buf)
    assert core.extract_stockmock_margin(buf.getvalue()) == pytest.approx(525000.0)
    assert core._parse_margin("2.66Cr") == pytest.approx(26600000.0)
    assert core._parse_margin("99,707") == pytest.approx(99707.0)
    assert core._parse_margin("not a number") is None

    t, _, _, _ = core.detect_and_load(wb, "basket.xlsx")
    dp = core.daily_pnl(t)
    assert dp["Trades"].sum() == len(t) and dp["Net P&L"].sum() == pytest.approx(t["P/L"].sum())
    assert (dp["Wins"] + dp["Losses"] <= dp["Trades"]).all()
